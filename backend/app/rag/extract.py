"""Turn an uploaded file into text, keeping PDF page numbers for citations."""

import csv
import io
import json
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePath


class ExtractionError(Exception):
    """Shown to the user next to the document."""


@dataclass(frozen=True)
class Page:
    number: int | None  # 1-based PDF page; None for formats without pages
    text: str


TEXT_TYPES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".java",
    ".go",
    ".sql",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".xml",
}
SUPPORTED = TEXT_TYPES | {".pdf", ".docx", ".csv", ".json", ".html", ".htm"}


def _decode(data: bytes) -> str:
    # UTF-16 only with its byte-order mark: almost any byte string "decodes" as
    # UTF-16, turning Latin-1/Windows text into garbage.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _TextOnly(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("This PDF is password-protected")
        pages = [Page(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    except PdfReadError as exc:
        raise ExtractionError("This PDF couldn't be read") from exc
    if not any(p.text.strip() for p in pages):
        raise ExtractionError("No text found. It may be a scanned PDF (images only).")
    return pages


def _docx(data: bytes) -> list[Page]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ExtractionError("This Word file couldn't be read") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [Page(None, "\n".join(parts))]


def extract(filename: str, data: bytes) -> list[Page]:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type {suffix or '(none)'}")
    if suffix == ".pdf":
        pages = _pdf(data)
    elif suffix == ".docx":
        pages = _docx(data)
    elif suffix == ".csv":
        rows = csv.reader(io.StringIO(_decode(data)))
        pages = [Page(None, "\n".join(", ".join(r) for r in rows))]
    elif suffix == ".json":
        try:
            text = json.dumps(json.loads(_decode(data)), indent=2, ensure_ascii=False)
        except ValueError:
            text = _decode(data)
        pages = [Page(None, text)]
    elif suffix in (".html", ".htm"):
        parser = _TextOnly()
        parser.feed(_decode(data))
        pages = [Page(None, "".join(parser.parts))]
    else:
        pages = [Page(None, _decode(data))]
    pages = [Page(p.number, re.sub(r"[ \t]+\n", "\n", p.text).strip()) for p in pages]
    if not any(p.text for p in pages):
        raise ExtractionError("The file has no text")
    return pages
